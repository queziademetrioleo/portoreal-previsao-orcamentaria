import io
import sys
import types
import unittest

sys.modules.setdefault('xlrd', types.ModuleType('xlrd'))
weasyprint = types.ModuleType('weasyprint')
weasyprint.HTML = object
sys.modules.setdefault('weasyprint', weasyprint)

from docx import Document

import relatorio_docx
import relatorio_pdf


def _estado(com_fundo=True):
    return {
        'nome_condominio': 'Aldeia das Asas', 'ano_previsao': 2027, 'com_fundo': com_fundo,
        'sistemas': ['group'], 'previsao_final': [],
        'linhas_contas': [
            {'grupo': 'Pessoal', 'classe': 'Salários', 'final': 12000},
            {'grupo': 'Contas de consumo', 'classe': 'Energia Elétrica', 'final': 2400},
            {'grupo': 'Zzz', 'classe': 'Conta sem categoria', 'final': 600},
        ],
        'fluxo_mensal': [{'mes': f'{m:02d}', 'receita': 2000, 'despesa': 1300, 'saldo': 700}
                         for m in range(1, 13)],
        'resumo': {
            'subtotal': 15000, 'total_previsto': 15000, 'inflacao': 0.10,
            'cenarios': {
                'fundo_reserva_anual': 3000,
                'com_fundo': {'receita_anual': 30000}, 'sem_fundo': {'receita_anual': 27000},
                'receitas_ordinarias': [{'classe': 'Taxas de Condomínio', 'mensal': 2250}],
                'receitas_nao_ordinarias': [],
            },
        },
    }


class RelatorioDocxTest(unittest.TestCase):
    def _textos(self, docx_bytes):
        doc = Document(io.BytesIO(docx_bytes))
        partes = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                partes += [c.text for c in row.cells]
        return '\n'.join(partes)

    def test_docx_tem_mesmas_secoes_do_pdf(self):
        texto = self._textos(relatorio_docx.gerar_relatorio_docx(_estado()))
        for trecho in ('Aldeia das Asas — 2027', 'Receitas', 'Despesas', 'Quadro de leitura',
                       'Conclusão', 'CONSIDERAÇÕES IMPORTANTES', 'TOTAL PREVISTO', 'PORTO REAL IMÓVEIS'):
            self.assertIn(trecho, texto)

    def test_sem_fundo_remove_fundo_nos_dois_formatos(self):
        ctx_sim = relatorio_pdf.montar_contexto_relatorio(_estado(), com_fundo_override=True)
        ctx_nao = relatorio_pdf.montar_contexto_relatorio(_estado(), com_fundo_override=False)
        self.assertTrue(any('fundo de reserva' in l.lower() for l, _ in ctx_sim['receitas']))
        self.assertFalse(any('fundo de reserva' in l.lower() for l, _ in ctx_nao['receitas']))
        self.assertFalse(any('Fundo de Reserva' in o['label'] for o in ctx_nao['opcoes']))
        self.assertFalse(any('Lembramos' in c for c in ctx_nao['consideracoes']))
        texto = self._textos(relatorio_docx.gerar_relatorio_docx(_estado(), com_fundo_override=False))
        self.assertNotIn('Taxa de Condomínio + Fundo de Reserva', texto)

    def test_outras_despesas_nao_aparece(self):
        ctx = relatorio_pdf.montar_contexto_relatorio(_estado())
        rotulos = [l for l, _ in ctx['despesas']]
        self.assertNotIn('Outras despesas', rotulos)
        self.assertNotIn('Ajustes de previsão', rotulos)
        # o resíduo (600) continua no total, dentro de conservação
        self.assertEqual(ctx['subtotal_mensal'], 'R$ 1.250,00')


if __name__ == '__main__':
    unittest.main()
