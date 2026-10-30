#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RELATORIO DOCX PARA O CONDOMINIO

Mesmo conteúdo e ordem do PDF (relatorio_pdf), editável no Word. Consome o
contexto de `montar_contexto_relatorio`, então toda regra do relatório
(fundo de reserva, despesas, opções) vale igual nos dois formatos.
O gráfico de barras entra como imagem PNG (desenhado com Pillow).
"""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from relatorio_pdf import (BARRA_CONTORNO, BARRA_COR, GRADE_FORTE, GRADE_FRACA,
                           _layout_barras, montar_contexto_relatorio)

AZUL = RGBColor(0x1C, 0x35, 0x5E)
CINZA = RGBColor(0x52, 0x51, 0x4E)
FONTE = 'Arial'
STATUS_TXT = {'superavit': 'Superávit', 'superavit_insuficiente': 'Superávit insuficiente',
              'deficit': 'Déficit'}
STATUS_COR = {'superavit': RGBColor(0x0F, 0x6B, 0x2B),
              'superavit_insuficiente': RGBColor(0x8A, 0x63, 0x00),
              'deficit': RGBColor(0xA1, 0x26, 0x26)}


def _fonte_img_negrito(tamanho):
    for caminho in ('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
                    '/System/Library/Fonts/Supplemental/Arial Bold.ttf'):
        try:
            return ImageFont.truetype(caminho, tamanho)
        except OSError:
            continue
    return _fonte_img(tamanho)


def _fonte_img(tamanho):
    for caminho in ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                    '/System/Library/Fonts/Supplemental/Arial.ttf'):
        try:
            return ImageFont.truetype(caminho, tamanho)
        except OSError:
            continue
    return ImageFont.load_default(size=tamanho)


def _hex(cor):
    cor = cor.lstrip('#')
    return tuple(int(cor[i:i + 2], 16) for i in (0, 2, 4))


def _png_barras(barras, escala=3):
    """Mesmo desenho do SVG do PDF (relatorio_pdf._layout_barras)."""
    lay = _layout_barras(barras)
    k = lambda v: round(v * escala)
    img = Image.new('RGB', (k(lay['largura']), k(lay['altura'])), 'white')
    d = ImageDraw.Draw(img)
    tinta = _hex('#14110c')
    fonte, fonte_eixo = _fonte_img(k(12.5)), _fonte_img(k(12))
    for g in lay['grade']:
        cor = _hex(GRADE_FORTE if g['forte'] else GRADE_FRACA)
        d.line([(k(g['x']), 0), (k(g['x']), k(lay['base']))], fill=cor, width=escala)
        if g['forte']:
            d.text((k(g['x']), k(lay['base'] + 18)), g['rotulo'], font=fonte_eixo, fill=tinta, anchor='mm')
    for l in lay['linhas']:
        if l['x_fim'] > lay['x0']:
            d.rectangle([k(lay['x0']), k(l['y']), k(l['x_fim']), k(l['y'] + lay['altura_barra'])],
                        fill=_hex(BARRA_COR), outline=_hex(BARRA_CONTORNO), width=max(1, k(1.2)))
        d.text((k(l['x_fim'] + 8), k(l['meio'])), l['pct'], font=_fonte_img(k(13)), fill=tinta, anchor='lm')
        n = len(l['rotulo'])
        for j, t in enumerate(l['rotulo']):
            d.text((k(lay['x0'] - 10), k(l['meio'] + (j - (n - 1) / 2) * 15)), t,
                   font=fonte, fill=tinta, anchor='rm')
    return _png_bytes(img)


def _png_bytes(img):
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def _sombrear(celula, hex_cor):
    pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_cor)
    pr.append(shd)


def _bordas(celula, topo=None):
    pr = celula._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{lado}')
        if lado == 'bottom':
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'E1E0D9')
        elif lado == 'top' and topo:
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(topo)); el.set(qn('w:color'), '14110C')
        else:
            el.set(qn('w:val'), 'nil')
        b.append(el)
    pr.append(b)


def _texto(par, txt, negrito=False, tamanho=10, cor=None):
    run = par.add_run(txt)
    run.bold = negrito
    run.font.size = Pt(tamanho)
    run.font.name = FONTE
    if cor is not None:
        run.font.color.rgb = cor
    return run


def _titulo(doc, txt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _texto(p, txt, True, 13, AZUL)
    pPr = p._p.get_or_add_pPr()
    borda = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'E1E0D9')
    borda.append(b)
    pPr.append(borda)


def _tabela_valores(doc, cab, linhas, destaques=()):
    """linhas = [(label, valor)]; destaques = {indice: 'total'|'subtotal'}."""
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    larguras = (Cm(12.4), Cm(4.6))
    for i, txt in enumerate(cab):
        c = t.rows[0].cells[i]
        c.width = larguras[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.LEFT
        _texto(p, txt, True, 10, CINZA)
        _bordas(c)
    for idx, (label, valor) in enumerate(linhas):
        tipo = dict(destaques).get(idx)
        cels = t.add_row().cells
        for i, txt in enumerate((label, valor)):
            cels[i].width = larguras[i]
            p = cels[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.LEFT
            _texto(p, txt, tipo is not None, 10, CINZA if tipo == 'subtotal' else None)
            _bordas(cels[i], topo=12 if tipo == 'total' else 6 if tipo == 'subtotal' else None)
    return t


def gerar_relatorio_docx(estado, logo_path=None, com_fundo_override=None):
    ctx = montar_contexto_relatorio(estado, logo_path, com_fundo_override)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    estilo = doc.styles['Normal']
    estilo.font.name = FONTE
    estilo.font.size = Pt(10)
    estilo.paragraph_format.space_after = Pt(0)

    # Rodapé: Página X de Y
    rodape = sec.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for pedaco in ('Página ', 'PAGE', ' de ', 'NUMPAGES'):
        if pedaco.isupper():
            run = rodape.add_run()
            run.font.size = Pt(9)
            for tipo, txt in (('begin', None), (None, pedaco), ('end', None)):
                if tipo:
                    el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tipo)
                else:
                    el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
                run._r.append(el)
        else:
            _texto(rodape, pedaco, False, 9, CINZA)

    # Cabeçalho: logo + título
    cab = doc.add_table(rows=1, cols=2)
    cab.autofit = False
    c_logo, c_tit = cab.rows[0].cells
    c_logo.width, c_tit.width = Cm(4), Cm(13)
    if logo_path:
        try:
            c_logo.paragraphs[0].add_run().add_picture(logo_path, height=Cm(1.8))
        except Exception:
            pass
    p = c_tit.paragraphs[0]
    _texto(p, f'{ctx["nome_condominio"]} — {ctx["ano_previsao"]}', True, 18, AZUL)
    p2 = c_tit.add_paragraph()
    _texto(p2, f'Previsão orçamentária · {ctx["data_extenso"]}', False, 11, CINZA)

    _titulo(doc, 'Receitas')
    _tabela_valores(doc, ('Conta', 'Valor mensal previsto'),
                    list(ctx['receitas']) + [('Total mensal', ctx['total_receitas'])],
                    {len(ctx['receitas']): 'total'})

    _titulo(doc, 'Despesas')
    desp = list(ctx['despesas'])
    n = len(desp)
    desp += [('SUBTOTAL', ctx['subtotal_mensal']),
             (f'Aumento Previsto (Salários, tarifas, serviços) = {ctx["inflacao_pct"]}% sobre o SUBTOTAL',
              ctx['aumento_mensal']),
             ('TOTAL PREVISTO', ctx['total_despesas'])]
    _tabela_valores(doc, ('Conta', 'Valor mensal previsto'), desp, {n: 'subtotal', n + 2: 'total'})

    _titulo(doc, 'Quadro de leitura')
    for o in ctx['opcoes']:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _texto(p, o['label'], True, 11, AZUL)
        t = _tabela_valores(doc, ('Item', 'Valor'), [
            ('Receita mensal', o['receita_mensal']),
            ('Despesa mensal', o['despesa_mensal']),
            ('Resultado mensal', o['resultado_mensal']),
            ('Reajuste necessário', o['reajuste']),
        ])
        c = t.add_row().cells
        c[0].merge(c[1])
        _texto(c[0].paragraphs[0], STATUS_TXT[o['status']], True, 10, STATUS_COR[o['status']])

    grafico = ctx.get('grafico')
    if grafico:
        _titulo(doc, 'Composição das despesas')
        doc.add_paragraph().add_run().add_picture(_png_barras(grafico['barras']), width=Cm(17))

    # Conclusão logo depois do gráfico, sem página nova, como no PDF.
    _titulo(doc, 'Conclusão')
    p = doc.add_paragraph()
    _texto(p, ctx['status_label'], True, 10, STATUS_COR[ctx['status_geral']])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    _texto(p, ctx['conclusao_texto'], False, 10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    _texto(p, 'CONSIDERAÇÕES IMPORTANTES', True, 13, AZUL)
    for i, texto in enumerate(ctx['consideracoes']):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        # Último item preso à assinatura: ela nunca fica sozinha numa página.
        p.paragraph_format.keep_with_next = i == len(ctx['consideracoes']) - 1
        _texto(p, texto, False, 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(24)
    _texto(p, f'Cabo Frio, {ctx["data_extenso"]}', False, 10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texto(p, 'PORTO REAL IMÓVEIS', True, 10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
