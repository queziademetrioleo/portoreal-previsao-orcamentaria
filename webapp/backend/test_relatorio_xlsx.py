import io
import sys
import types
import unittest

sys.modules.setdefault('xlrd', types.ModuleType('xlrd'))
weasyprint = types.ModuleType('weasyprint')
weasyprint.HTML = object
sys.modules.setdefault('weasyprint', weasyprint)

from openpyxl import load_workbook

import relatorio_pdf
import relatorio_xlsx
from test_relatorio_docx import _estado


def _linha(ws, texto):
    return next(r for r in range(1, ws.max_row + 1) if ws[f'B{r}'].value == texto)


def _todas(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


class RelatorioXlsxTest(unittest.TestCase):
    def setUp(self):
        self.bytes = relatorio_xlsx.gerar_relatorio_xlsx(_estado())
        self.ws = load_workbook(io.BytesIO(self.bytes))['Previsão']

    def test_mesmas_secoes_do_pdf(self):
        textos = _todas(self.ws)
        for trecho in ('Aldeia das Asas — 2027', 'Receitas', 'Despesas', 'Quadro de leitura',
                       'Composição das despesas', 'Conclusão',
                       'CONSIDERAÇÕES IMPORTANTES', 'TOTAL PREVISTO', 'PORTO REAL IMÓVEIS'):
            self.assertIn(trecho, textos)

    def test_totais_sao_formulas_com_linha_livre(self):
        ws = self.ws
        tot_rec = _linha(ws, 'Total mensal')
        self.assertTrue(ws[f'F{tot_rec}'].value.startswith('=SUM(F'))
        # Linha em branco dentro do intervalo somado, para incluir conta nova.
        self.assertIsNone(ws[f'B{tot_rec - 1}'].value)
        sub = _linha(ws, 'SUBTOTAL')
        self.assertIsNone(ws[f'B{sub - 1}'].value)
        self.assertEqual(ws[f'F{sub + 1}'].value, f'=ROUND(F{sub}*E{sub + 1},2)')
        self.assertAlmostEqual(ws[f'E{sub + 1}'].value, 0.10)
        self.assertEqual(ws[f'F{sub + 2}'].value, f'=F{sub}+F{sub + 1}')

    def test_fundo_separado_nas_opcoes(self):
        ws = self.ws
        sem = _linha(ws, 'Receitas sem fundo de reserva')
        self.assertIn('-SUMIF(', ws[f'C{sem + 1}'].value)
        self.assertNotIn('+SUMIF(', ws[f'C{sem + 1}'].value)
        self.assertIn('+SUMIF(', ws[f'F{sem + 1}'].value)

    def test_formulas_levam_resultado_em_cache(self):
        # Prévias (Finder, WhatsApp) não calculam: o valor tem de vir gravado.
        cache = load_workbook(io.BytesIO(self.bytes), data_only=True)['Previsão']
        tot_rec = _linha(self.ws, 'Total mensal')
        self.assertAlmostEqual(cache[f'F{tot_rec}'].value, 2250 + 250)
        self.assertTrue(any('prática não recomendada' in t
                            for t in _todas(cache)))

    def test_conclusao_sem_resultado_mensal(self):
        self.assertFalse(any('O resultado mensal é de' in t for t in _todas(self.ws)))

    def test_texto_do_reajuste_igual_ao_pdf(self):
        opcoes = [{'label': 'sem', 'usa_fundo': False}, {'label': 'com', 'usa_fundo': True}]
        for res_sem, reaj_sem, res_com, reaj_com in ((-1000.0, 0.13, -500.0, 0.10), (-1000.0, 0.13, 200.0, 0.0),
                                                     (300.0, 0.0, 800.0, 0.0)):
            refs = [{'res': 'F1', 'reaj': 'F2', 'resultado': res_sem, 'reajuste': reaj_sem},
                    {'res': 'F3', 'reaj': 'F4', 'resultado': res_com, 'reajuste': reaj_com}]
            formula, texto = relatorio_xlsx._texto_opcoes(opcoes, refs)
            pdf = relatorio_pdf._texto_consideracao(relatorio_pdf._consideracao_opcoes(
                [dict(o, reajuste=r['reajuste']) for o, r in zip(opcoes, refs)]))
            self.assertEqual(texto, pdf)
            self.assertIn('FIXED(F4*100,1)', formula)

    def test_grafico_de_pizza(self):
        self.assertEqual(len(self.ws._charts), 1)
